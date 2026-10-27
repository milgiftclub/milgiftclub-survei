import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
import os
from questions import Q, HIGH, RANDOMIZE, RANDOMIZE_ROWS, variant_v2
VAR=os.environ.get('V','v1')
if VAR=='v2': Q=variant_v2()
from categories import CATS
A='Arial'; YEL='FFFF00'
thin=Side(style='thin',color='BFBFBF'); B=Border(top=thin,bottom=thin,left=thin,right=thin)
num={q[0]:i+1 for i,q in enumerate(Q)}
def res(s): return re.sub(r'\{(\w+)\}', lambda m: f'Q{num[m.group(1)]}', s)

wb=Workbook(); ws=wb.active; ws.title='Questionnaire'
hdr=['No','Topic','Pertanyaan','Tipe','Attributes','Routing / Skip Logic','Notes','Notes from MIL Team','Preview']
for j,h in enumerate(hdr,1):
    c=ws.cell(1,j,h); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL)
    c.alignment=Alignment(horizontal='center',vertical='center'); c.border=B
base=Font(name=A,size=9); top=Alignment(vertical='top',wrap_text=True)
rows_by_topic=[]
for i,(qid,topic,qt,typ,attrs,rout,notes) in enumerate(Q):
    r=i+2; n=i+1
    qt=res(qt); attrs=[res(a) for a in attrs]; rout=res(rout).strip(); notes=res(notes)
    _add=[]
    if qid in RANDOMIZE: _add.append('Urutan opsi diacak per responden (opsi Lainnya / Tidak ... / eksklusif tetap di akhir)')
    if qid in RANDOMIZE_ROWS: _add.append('Urutan baris pernyataan diacak per responden')
    if qid=='C1': _add.append('Opsi tidak diacak; tampil keterangan "Gulir ke bawah untuk melihat kategori lainnya"')
    if _add: rout='\n'.join(([rout] if rout not in ('-','') else [])+_add)
    preview=f'{n}. {qt}\n' + '\n'.join(attrs)
    vals=[n,topic,qt,typ,'\n'.join(attrs),rout,notes,'',preview]
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=base; c.border=B; c.alignment=top
    ws.cell(r,1).font=Font(name=A,size=9,bold=True)
    ws.cell(r,4).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    lines=max(len(attrs)+1, len(qt)//45+1, len(rout)//40+1, len(notes)//45+1)
    ws.row_dimensions[r].height=min(max(lines*12+6,30),409)
    if rows_by_topic and rows_by_topic[-1][0]==topic: rows_by_topic[-1][2]=r
    else: rows_by_topic.append([topic,r,r])
for topic,r1,r2 in rows_by_topic:
    if r2>r1: ws.merge_cells(start_row=r1,start_column=2,end_row=r2,end_column=2)
    c=ws.cell(r1,2); c.font=Font(name=A,size=9,bold=True); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for k,v in {'A':5,'B':18,'C':45,'D':11,'E':42,'F':38,'G':42,'H':28,'I':55}.items(): ws.column_dimensions[k].width=v
ws.freeze_panes='C2'

# --- Daftar Kategori
w2=wb.create_sheet('Daftar Kategori')
for j,h in enumerate(['No','Kelompok','Kategori','Contoh produk','Prioritas','Masuk modul funnel?'],1):
    c=w2.cell(1,j,h); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL); c.border=B; c.alignment=Alignment(horizontal='center')
for i,c in enumerate(CATS):
    r=i+2; p='High' if c[1] in HIGH else 'Low'
    vals=[i+1,c[0],c[1],c[2],p,f'=IF(E{r}="High","Ya - loop funnel","Tidak - penetrasi saja")']
    for j,v in enumerate(vals,1):
        x=w2.cell(r,j,v); x.font=Font(name=A,size=9,bold=(j==3 and p=='High'),color='0000FF' if j==5 else '000000'); x.border=B; x.alignment=top
dv=DataValidation(type='list',formula1='"High,Low"'); w2.add_data_validation(dv); dv.add(f'E2:E{len(CATS)+1}')
n=len(CATS)+3
w2.cell(n,2,'Jumlah kategori High').font=Font(name=A,size=9,bold=True)
w2.cell(n,5,f'=COUNTIF(E2:E{len(CATS)+1},"High")').font=Font(name=A,size=9,bold=True)
w2.cell(n+1,2,'Kolom Prioritas (biru) dapat diubah; modul funnel di sheet Questionnaire mengikuti kategori High.').font=Font(name=A,size=8,italic=True)
for k,v in {'A':5,'B':22,'C':26,'D':38,'E':10,'F':22}.items(): w2.column_dimensions[k].width=v
w2.freeze_panes='A2'

# --- Daftar Produk (PT1) & rentang harga (PT2)
from products import PRODUCTS, PRICE_BANDS
wp=wb.create_sheet('Daftar Produk')
for j,h in enumerate(['Kategori','Pilihan jenis produk (Q PT1)'],1):
    c=wp.cell(1,j,h); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL); c.border=B
rr=2
for cat,lst in PRODUCTS.items():
    if VAR=='v2' and cat not in HIGH: continue
    x=wp.cell(rr,1,cat); x.font=Font(name=A,size=9,bold=True); x.border=B; x.alignment=top
    y=wp.cell(rr,2,'\n'.join(lst)); y.font=Font(name=A,size=9); y.border=B; y.alignment=top; rr+=1
wp.cell(rr+1,1,'Rentang harga per item (PT2)').font=Font(name=A,size=9,bold=True)
wp.cell(rr+1,2,'\n'.join(PRICE_BANDS)).font=Font(name=A,size=9); wp.cell(rr+1,2).alignment=top
for k,v in {'A':30,'B':60}.items(): wp.column_dimensions[k].width=v
wp.freeze_panes='A2'

# --- Daftar Kota
import json as _j
wk=wb.create_sheet('Daftar Kota')
for j,h in enumerate(['No','Kabupaten/Kota','Provinsi','Tipe wilayah','Jabodetabek'],1):
    c=wk.cell(1,j,h); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL); c.border=B
for i,cty in enumerate(_j.load(open('cities.json'))):
    for j,v in enumerate([i+1]+cty,1):
        x=wk.cell(i+2,j,v); x.font=Font(name=A,size=9); x.border=B
wk.cell(len(_j.load(open('cities.json')))+3,2,'Sumber: Kepmendagri No 300.2.2-2138 Tahun 2025 (via github.com/cahyadsn/wilayah). Provinsi, tipe wilayah, dan Jabodetabek diturunkan otomatis dari jawaban kota.').font=Font(name=A,size=8,italic=True)
for k,v in {'A':6,'B':32,'C':26,'D':14,'E':12}.items(): wk.column_dimensions[k].width=v
wk.freeze_panes='A2'

# --- Estimasi Durasi
w3=wb.create_sheet('Estimasi Durasi')
w3['A1']='Estimasi durasi pengisian survei'; w3['A1'].font=Font(name=A,bold=True,size=12)
w3['A2']='Input (teks biru, sel kuning) dapat diubah. Asumsi detik per tipe pertanyaan berdasarkan pengisian di ponsel; validasi dengan hasil pilot (n≈30).'
w3['A2'].font=Font(name=A,size=8,italic=True)
types=[('SA',8),('SA (dropdown)',10),('SA (cari kota)',12),('MA',12),('MA (maks. 3)',14),('Matrix SA',25),('Matrix MA',35),('NPS 0-10',8),('OA',25),('Info',5),('Hidden',0)]
w3['A4']='Tipe'; w3['B4']='Detik / pertanyaan'
for j in (1,2): c=w3.cell(4,j); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL); c.border=B
for i,(t,s) in enumerate(types):
    r=5+i; w3.cell(r,1,t).font=base; c=w3.cell(r,2,s); c.font=Font(name=A,size=9,color='0000FF'); c.fill=PatternFill('solid',fgColor='FFF4CC')
    for j in (1,2): w3.cell(r,j).border=B
tr0,tr1=5,4+len(types)
w3['D4']='Jumlah kategori modul funnel per responden'; w3['D4'].font=Font(name=A,bold=True,size=9)
c=w3['G4']; c.value=1; c.font=Font(name=A,size=9,color='0000FF',bold=True); c.fill=PatternFill('solid',fgColor='FFF4CC'); c.border=B
w3['D5']='Porsi tampil pertanyaan bersyarat (skip logic)'; w3['D5'].font=Font(name=A,bold=True,size=9)
c=w3['G5']; c.value=0.6; c.number_format='0%'; c.font=Font(name=A,size=9,color='0000FF',bold=True); c.fill=PatternFill('solid',fgColor='FFF4CC'); c.border=B
w3['H5']='Asumsi: pertanyaan dengan "Tampilkan HANYA jika" muncul untuk ±60% responden'; w3['H5'].font=Font(name=A,size=8,italic=True)
h0=17
for j,h in enumerate(['No','Topic','Tipe','Modul kategori?','Bersyarat?','Detik dasar','Detik efektif'],1):
    c=w3.cell(h0,j,h); c.font=Font(name=A,bold=True,size=9); c.fill=PatternFill('solid',fgColor=YEL); c.border=B
for i,(qid,topic,qt,typ,attrs,rout,notes) in enumerate(Q):
    r=h0+1+i
    loop='Ya' if 'MODUL KATEGORI' in rout else 'Tidak'
    cond='Ya' if 'HANYA jika' in rout else 'Tidak'
    vals=[i+1,topic.split('\n')[0],typ,loop,cond,
          f'=IFERROR(INDEX($B${tr0}:$B${tr1},MATCH(C{r},$A${tr0}:$A${tr1},0)),0)',
          f'=F{r}*IF(D{r}="Ya",$G$4,1)*IF(E{r}="Ya",$G$5,1)']
    for j,v in enumerate(vals,1):
        c=w3.cell(r,j,v); c.font=base; c.border=B
last=h0+len(Q)
w3['D7']='Total estimasi (menit)'; w3['D7'].font=Font(name=A,bold=True,size=10)
w3['G7']=f'=ROUND(SUM(G{h0+1}:G{last})/60,1)'; w3['G7'].font=Font(name=A,bold=True,size=11)
w3['D8']='Target maksimal (menit)'; w3['D8'].font=Font(name=A,bold=True,size=9); w3['G8']=15; w3['G8'].font=Font(name=A,size=9,color='0000FF')
w3['D9']='Status'; w3['D9'].font=Font(name=A,bold=True,size=9)
w3['G9']='=IF(G7<=G8,"Sesuai target",IF(G7<=G8*1.2,"Sedikit di atas target","Terlalu panjang - perlu dipangkas"))'; w3['G9'].font=Font(name=A,bold=True,size=9)
w3['D11']='Durasi bagian funnel (menit, termasuk loop)'; w3['D11'].font=Font(name=A,size=9)
w3['G11']=f'=ROUND(SUMIF(D{h0+1}:D{last},"Ya",G{h0+1}:G{last})/60,1)'; w3['G11'].font=Font(name=A,size=9)
w3['D12']='Jumlah pertanyaan tampil (tanpa hidden & info)'; w3['D12'].font=Font(name=A,size=9)
w3['G12']=f'=COUNTA(C{h0+1}:C{last})-COUNTIF(C{h0+1}:C{last},"Hidden")-COUNTIF(C{h0+1}:C{last},"Info")'; w3['G12'].font=Font(name=A,size=9)
for k,v in {'A':14,'B':26,'C':14,'D':10,'E':11,'F':11,'G':12,'H':40}.items(): w3.column_dimensions[k].width=v
wb.save('Kuesioner_v1_SemuaKategori_v29.xlsx' if VAR=='v1' else 'Kuesioner_v2_FashionTasHijab_v29.xlsx')
print(len(Q))
